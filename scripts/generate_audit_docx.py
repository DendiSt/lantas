import os
import docx
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_ALIGN_VERTICAL
from docx.oxml import parse_xml, OxmlElement
from docx.oxml.ns import nsdecls, qn

def set_cell_background(cell, fill_hex):
    tcPr = cell._element.get_or_add_tcPr()
    shd = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{fill_hex}"/>')
    tcPr.append(shd)

def set_cell_margins(cell, top=100, bottom=100, left=150, right=150):
    tcPr = cell._element.get_or_add_tcPr()
    tcMar = parse_xml(f'''
        <w:tcMar {nsdecls("w")}>
            <w:top w:w="{top}" w:type="dxa"/>
            <w:bottom w:w="{bottom}" w:type="dxa"/>
            <w:left w:w="{left}" w:type="dxa"/>
            <w:right w:w="{right}" w:type="dxa"/>
        </w:tcMar>
    ''')
    tcPr.append(tcMar)

def create_audit_document(filename="AUDIT_CHECKLIST_LANTAS.docx"):
    doc = Document()

    # Set page margins to 1 inch
    sections = doc.sections
    for section in sections:
        section.top_margin = Inches(1.0)
        section.bottom_margin = Inches(1.0)
        section.left_margin = Inches(1.0)
        section.right_margin = Inches(1.0)

    # Base Styles
    normal_style = doc.styles['Normal']
    normal_style.font.name = 'Calibri'
    normal_style.font.size = Pt(11)
    normal_style.font.color.rgb = RGBColor(30, 41, 59) # Slate 800

    # Header / Title Block
    title_p = doc.add_paragraph()
    title_p.paragraph_format.space_before = Pt(0)
    title_p.paragraph_format.space_after = Pt(2)
    run_badge = title_p.add_run("LAPORAN AUDIT RESMI & CHECKLIST VERIFIKASI")
    run_badge.font.size = Pt(9.5)
    run_badge.font.bold = True
    run_badge.font.color.rgb = RGBColor(15, 118, 110) # Teal 700

    title = doc.add_paragraph()
    title.paragraph_format.space_before = Pt(0)
    title.paragraph_format.space_after = Pt(4)
    run_title = title.add_run("LANTAS — Layanan Terpadu Administrasi Sekolah")
    run_title.font.size = Pt(22)
    run_title.font.bold = True
    run_title.font.color.rgb = RGBColor(15, 23, 42) # Slate 900

    sub = doc.add_paragraph()
    sub.paragraph_format.space_before = Pt(0)
    sub.paragraph_format.space_after = Pt(14)
    run_sub = sub.add_run("Evaluasi Kesiapan MVP Prototype & Deliverables • SMK Negeri 2 Subang")
    run_sub.font.size = Pt(12)
    run_sub.font.italic = True
    run_sub.font.color.rgb = RGBColor(100, 116, 139) # Slate 500

    # Metadata Callout Box
    meta_table = doc.add_table(rows=1, cols=1)
    meta_table.alignment = WD_TABLE_ALIGNMENT.CENTER
    meta_cell = meta_table.cell(0, 0)
    set_cell_background(meta_cell, "F8FAFC") # Slate 50
    set_cell_margins(meta_cell, top=140, bottom=140, left=200, right=200)
    
    mp = meta_cell.paragraphs[0]
    mp.paragraph_format.space_after = Pt(2)
    m_run1 = mp.add_run("📋 Ringkasan Audit: ")
    m_run1.bold = True
    mp.add_run("Audit komprehensif terhadap fitur fungsional, arsitektur teknis, kelengkapan deliverable, dan repositori proyek LANTAS berdasarkan standar Feature Completion & Prototype Deliverables.")
    
    mp2 = meta_cell.add_paragraph()
    mp2.paragraph_format.space_before = Pt(4)
    mp2.paragraph_format.space_after = Pt(0)
    mp2.add_run("• Repositori: ").bold = True
    mp2.add_run("https://github.com/DendiSt/lantas.git  |  ")
    mp2.add_run("• Stack: ").bold = True
    mp2.add_run("Next.js 16 (App Router), Prisma ORM, SQLite, Tailwind CSS")

    doc.add_paragraph().paragraph_format.space_after = Pt(12)

    # Helper function for adding sections
    def add_section_header(text, icon="📌"):
        h = doc.add_paragraph()
        h.paragraph_format.space_before = Pt(16)
        h.paragraph_format.space_after = Pt(8)
        run = h.add_run(f"{icon} {text}")
        run.font.size = Pt(14)
        run.font.bold = True
        run.font.color.rgb = RGBColor(15, 23, 42)

    # 1. Feature Completion Section
    add_section_header("1. Checklist Feature Completion", "✅")
    
    desc1 = doc.add_paragraph("Sebuah fitur dianggap selesai (Done) apabila memenuhi seluruh indikator berikut:")
    desc1.paragraph_format.space_after = Pt(6)

    headers1 = ["Kriteria / Indikator", "Status", "Bukti Implementasi & File Rujukan", "Catatan Hasil Evaluasi"]
    col_widths1 = [Inches(1.8), Inches(0.9), Inches(2.2), Inches(2.1)]
    
    data1 = [
        ("Requirement sudah jelas", "[x] Lolos", "PROJECT_CONTEXT.md (Line 73-125)", "Product Goal, Persona Siswa & Staf TU, Core Problems, dan spesifikasi MVP F-001 s/d F-004 terdokumentasi terinci."),
        ("UI sudah sesuai wireframe", "[x] Lolos", "src/app/dashboard/page.tsx\nsrc/app/admin/page.tsx", "Palet warna monokromatik netral (Slate/Zinc), tombol dark charcoal, soft status badges (Amber/Emerald/Rose), serta layout desktop & mobile terpisah."),
        ("Functionality berjalan", "[x] Lolos", "src/app/actions/requests.ts\nprisma/schema.prisma", "Siswa berhasil mengajukan izin, upload bukti lampiran, dan TU dapat memverifikasi Approve / Reject secara instan."),
        ("Validation sudah dilakukan", "[x] Lolos", "CreateRequestDialog.tsx\nsrc/app/actions/requests.ts", "Validasi form di browser (jenis izin, alasan wajib diisi) serta validasi sisi server sebelum eksekusi database."),
        ("Error state sudah ditangani", "[x] Lolos", "next.config.ts\nAdminRequestsTable.tsx\nRequestHistoryList.tsx", "Batas payload Server Action dinaikkan (10MB) + kompresi kanvas; penanganan empty state jika daftar pengajuan kosong."),
        ("Responsive jika diperlukan", "[x] Lolos", "BottomNav.tsx\nAdminSidebar.tsx\ndashboard/page.tsx", "Portal Siswa mobile-first dengan floating bottom navigation bar; Panel TU memiliki responsive sidebar & tabel scroll aman."),
        ("Tidak ada console error", "[x] Lolos", "npx tsc --noEmit (0 error)\nnpm run lint (0 warning)", "Lolos audit kompilasi TypeScript murni, zero warning ESLint, dan build produksi Turbopack sukses tanpa kendala."),
        ("Sudah di-test", "[x] Lolos", "Automated Browser Automation &\nManual Flow Check", "Pengujian E2E mencakup simulasi siswa input izin, kompresi gambar, approval real-time, serta filter status riwayat."),
        ("Code sudah direview", "[x] Lolos", "Direktori src/\nsrc/lib/prisma.ts", "Pemisahan Server/Client Component sudah tepat, implementasi Singleton Prisma bersih, dan penghapusan implicit-any."),
        ("PR sudah di-merge", "[ ] Pending", "Branch: origin/main\n(Belum ada cabang PR)", "Pengembangan saat ini di-push langsung ke cabang main. Belum ada rekam jejak Pull Request antar-branch di GitHub."),
        ("Dokumentasi diperbarui", "[x] Lolos", "README.md\nPROJECT_CONTEXT.md", "README memuat dokumentasi resmi SMK Negeri 2 Subang, cara install, seeding, dan panduan fitur lengkap.")
    ]

    table1 = doc.add_table(rows=1, cols=4)
    table1.alignment = WD_TABLE_ALIGNMENT.CENTER
    table1.autofit = False

    hdr_cells = table1.rows[0].cells
    for i, title_text in enumerate(headers1):
        cell = hdr_cells[i]
        cell.width = col_widths1[i]
        set_cell_background(cell, "0F172A") # Slate 900
        set_cell_margins(cell, top=120, bottom=120, left=140, right=140)
        p = cell.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.LEFT
        run = p.add_run(title_text)
        run.bold = True
        run.font.size = Pt(9.5)
        run.font.color.rgb = RGBColor(255, 255, 255)

    for row_idx, row_data in enumerate(data1):
        row = table1.add_row()
        bg_color = "FFFFFF" if row_idx % 2 == 0 else "F8FAFC"
        for i, text in enumerate(row_data):
            cell = row.cells[i]
            cell.width = col_widths1[i]
            set_cell_background(cell, bg_color)
            set_cell_margins(cell, top=100, bottom=100, left=140, right=140)
            p = cell.paragraphs[0]
            p.paragraph_format.space_after = Pt(0)
            
            run = p.add_run(text)
            run.font.size = Pt(9)
            if i == 1:
                if "[x]" in text:
                    run.bold = True
                    run.font.color.rgb = RGBColor(5, 150, 105) # Emerald 600
                else:
                    run.bold = True
                    run.font.color.rgb = RGBColor(225, 29, 72) # Rose 600

    # 2. Final Prototype Deliverables Section
    doc.add_page_break()
    add_section_header("2. Checklist Final Prototype Deliverables", "📦")
    doc.add_paragraph("Target capaian akhir fase prototyping yang wajib dimiliki tim:").paragraph_format.space_after = Pt(6)

    # Category A: Product
    sub1 = doc.add_paragraph()
    sub1.paragraph_format.space_before = Pt(6)
    sub1.paragraph_format.space_after = Pt(4)
    r1 = sub1.add_run("A. Product Deliverables")
    r1.bold = True
    r1.font.size = Pt(11)
    r1.font.color.rgb = RGBColor(30, 41, 59)

    data_prod = [
        ("Working prototype", "[x] Lolos", "Aplikasi berjalan aktif di http://localhost:3000 dengan Next.js 16 App Router & Turbopack."),
        ("Core user flow berjalan", "[x] Lolos", "Alur Siswa ajukan izin -> TU verifikasi/tolak -> Siswa cek status berjalan real-time tanpa refresh."),
        ("UI konsisten", "[x] Lolos", "Desain terpadu berbasis Clean Minimalist Monochromatic Tailwind CSS dan shadcn/ui tokens."),
        ("Responsive", "[x] Lolos", "Adaptif sempurna di viewport mobile (375px–430px) hingga layar lebar desktop (>1280px).")
    ]

    for item, status, desc in data_prod:
        p = doc.add_paragraph()
        p.paragraph_format.left_indent = Inches(0.2)
        p.paragraph_format.space_after = Pt(3)
        r_status = p.add_run(f"{status}  ")
        r_status.bold = True
        r_status.font.color.rgb = RGBColor(5, 150, 105)
        r_item = p.add_run(f"{item}: ")
        r_item.bold = True
        p.add_run(desc)

    # Category B: Documentation
    sub2 = doc.add_paragraph()
    sub2.paragraph_format.space_before = Pt(10)
    sub2.paragraph_format.space_after = Pt(4)
    r2 = sub2.add_run("B. Documentation Deliverables")
    r2.bold = True
    r2.font.size = Pt(11)
    r2.font.color.rgb = RGBColor(30, 41, 59)

    data_doc = [
        ("Product Plan", "[ ] Parsial", "Tercakup di PROJECT_CONTEXT.md (Product Goal & Objective), belum dibuatkan berkas dokumen individual."),
        ("PRD (Product Requirements Document)", "[ ] Parsial", "Tercakup di PROJECT_CONTEXT.md (Problem Discovery & Requirement Summary), belum berupa file tersendiri."),
        ("Product Specification", "[ ] Parsial", "Spesifikasi fitur dan batasan sistem tercatat di PROJECT_CONTEXT.md."),
        ("User Stories", "[ ] Parsial", "Tercantum pemetaan US-001 (Siswa Ajukan Izin), US-002 (Riwayat), US-003 (Verifikasi TU)."),
        ("User Flow", "[ ] Parsial", "Alur operasional tertulis di PROJECT_CONTEXT.md (Line 88-91), belum berupa bagan/diagram visual."),
        ("Feature Specification", "[ ] Parsial", "Spesifikasi detail F-001 s/d F-004 tertulis lengkap di Project Context."),
        ("MVP Scope", "[ ] Parsial", "Batasan MVP (eksklusi integrasi WA & auth kompleks) terdefinisi jelas di Project Context."),
        ("Technical Specification", "[ ] Parsial", "Arsitektur Next.js Server Actions, Prisma SQLite, dan struktur folder terdokumentasi."),
        ("Data Model", "[x] Lolos", "Terdefinisi lengkap dan hidup di prisma/schema.prisma (Model User & Request dengan enums)."),
        ("Release Plan", "[ ] Belum Ada", "Belum disusun dokumen tahapan rilis formal dari MVP menuju deployment sekolah.")
    ]

    for item, status, desc in data_doc:
        p = doc.add_paragraph()
        p.paragraph_format.left_indent = Inches(0.2)
        p.paragraph_format.space_after = Pt(3)
        r_status = p.add_run(f"{status}  ")
        r_status.bold = True
        if "[x]" in status:
            r_status.font.color.rgb = RGBColor(5, 150, 105)
        else:
            r_status.font.color.rgb = RGBColor(217, 119, 6) # Amber 600
        r_item = p.add_run(f"{item}: ")
        r_item.bold = True
        p.add_run(desc)

    # Category C: Engineering
    sub3 = doc.add_paragraph()
    sub3.paragraph_format.space_before = Pt(10)
    sub3.paragraph_format.space_after = Pt(4)
    r3 = sub3.add_run("C. Engineering Deliverables")
    r3.bold = True
    r3.font.size = Pt(11)
    r3.font.color.rgb = RGBColor(30, 41, 59)

    data_eng = [
        ("GitHub Repository", "[x] Lolos", "Terhubung resmi ke repositori publik: https://github.com/DendiSt/lantas.git"),
        ("GitHub Project", "[ ] Belum Ada", "Belum dibuat papan Project/Kanban board di GitHub Projects web interface."),
        ("Clean branch structure", "[ ] Belum Ada", "Baru memiliki 1 branch tunggal (main), belum ada branch development atau feature-branch."),
        ("Meaningful commit history", "[x] Lolos", "Commit history rapi dan bermakna (contoh: 'feat: complete MVP prototype implementation')."),
        ("Pull Requests", "[ ] Belum Ada", "Belum ada riwayat PR di GitHub karena proses commit dilakukan langsung ke main."),
        ("README", "[x] Lolos", "README.md sangat komprehensif, mencakup konteks SMKN 2 Subang, arsitektur, dan cara instalasi."),
        ("Working prototype", "[x] Lolos", "Aplikasi siap diuji dan bebas dari build/runtime error.")
    ]

    for item, status, desc in data_eng:
        p = doc.add_paragraph()
        p.paragraph_format.left_indent = Inches(0.2)
        p.paragraph_format.space_after = Pt(3)
        r_status = p.add_run(f"{status}  ")
        r_status.bold = True
        if "[x]" in status:
            r_status.font.color.rgb = RGBColor(5, 150, 105)
        else:
            r_status.font.color.rgb = RGBColor(225, 29, 72)
        r_item = p.add_run(f"{item}: ")
        r_item.bold = True
        p.add_run(desc)

    # 3. Action Plan / Remediation Section
    doc.add_page_break()
    add_section_header("3. Rencana Tindak Lanjut & Solusi Perbaikan", "🎯")

    act_intro = doc.add_paragraph("Untuk menyempurnakan seluruh item yang masih berstatus [ ] Pending atau [ ] Parsial agar siap 100% saat penilaian akhir, berikut rencana aksi yang disarankan:")
    act_intro.paragraph_format.space_after = Pt(8)

    actions = [
        ("Langkah 1: Membuat Paket Berkas Dokumentasi Mandiri (Folder docs/)", 
         "Pecah ringkasan yang ada di PROJECT_CONTEXT.md menjadi file dokumen formal mandiri: Product Plan (docs/01-product-plan.md), PRD (docs/02-prd.md), User Stories & Flow (docs/03-user-stories.md), Technical Spec (docs/04-tech-spec.md), dan Release Plan (docs/05-release-plan.md)."),
        
        ("Langkah 2: Memenuhi Kriteria Pull Request (PR) di GitHub", 
         "1. Buat branch baru dari terminal lokal: git checkout -b docs/prototype-specifications\n"
         "2. Tambahkan berkas dokumentasi di atas, lakukan commit: git commit -m 'docs: add complete PRD and specifications'\n"
         "3. Push cabang baru: git push origin docs/prototype-specifications\n"
         "4. Buka halaman GitHub, klik 'Compare & pull request', lalu lakukan 'Merge Pull Request' ke main. Hal ini otomatis memenuhi kriteria 'Pull Requests' dan 'PR sudah di-merge'."),
        
        ("Langkah 3: Membuat GitHub Project Board", 
         "Buka https://github.com/DendiSt/lantas/projects, klik 'New Project', pilih template 'Kanban Board'. Buat 4 kartu backlog berdasarkan fitur LANTAS: F-001 (Formulir Izin), F-002 (Riwayat Izin), F-003 (Tabel Masuk TU), dan F-004 (Aksi Approve/Reject).")
    ]

    for act_title, act_desc in actions:
        ap = doc.add_paragraph()
        ap.paragraph_format.space_before = Pt(6)
        ap.paragraph_format.space_after = Pt(2)
        r_at = ap.add_run(act_title)
        r_at.bold = True
        r_at.font.size = Pt(10.5)
        r_at.font.color.rgb = RGBColor(15, 23, 42)

        dp = doc.add_paragraph(act_desc)
        dp.paragraph_format.left_indent = Inches(0.2)
        dp.paragraph_format.space_after = Pt(6)
        dp.style.font.size = Pt(9.5)
        dp.style.font.color.rgb = RGBColor(71, 85, 105)

    # Footer note
    doc.add_paragraph().paragraph_format.space_after = Pt(16)
    footer_p = doc.add_paragraph("Dokumen ini digenerate secara otomatis untuk audit kesiapan proyek LANTAS SMK Negeri 2 Subang.")
    footer_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    footer_p.style.font.size = Pt(8.5)
    footer_p.style.font.color.rgb = RGBColor(148, 163, 184)

    doc.save(filename)
    print(f"[SUCCESS] Dokumen berhasil dibuat: {filename}")

if __name__ == "__main__":
    create_audit_document("AUDIT_CHECKLIST_LANTAS.docx")
